"""Builds Docs/XR-Continuum-Viva-Presentation.pptx from the EaTemp coffee template.

Run from BaristaVR/:  python Docs/tools/make_viva_deck.py
Screenshots come from Docs/screenshots/ (see extract_screenshots.py).
Template slides are cloned only for their decorative art; all content is native, editable shapes.
"""
import copy
import os
import re
import sys

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

DOCS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(DOCS, "Coffee PowerPoint Template by EaTemp.pptx")
BARISTA = os.path.join(DOCS, "Untitled design (1).png")
SHOTS = os.path.join(DOCS, "screenshots")
OUT = os.path.join(DOCS, "XR-Continuum-Viva-Presentation.pptx")
TMP = os.path.join(os.environ.get("TEMP", DOCS), "viva_deck_tmp")
os.makedirs(TMP, exist_ok=True)

FONT = "Google Sans"
FONT_SEMI = "Google Sans SemiBold"
FONT_MED = "Google Sans Medium"

CREAM = RGBColor(0xE9, 0xD7, 0xC2)
PAPER = RGBColor(0xF6, 0xEE, 0xE3)
PAPER_ALT = RGBColor(0xF0, 0xE4, 0xD5)
RULE = RGBColor(0xD6, 0xC0, 0xA8)
DARK = RGBColor(0x3D, 0x26, 0x1A)
DARK_CARD = RGBColor(0x5E, 0x3E, 0x2E)
DARK_ALT = RGBColor(0x53, 0x36, 0x28)
DARK_RULE = RGBColor(0x7D, 0x5A, 0x46)
INK = RGBColor(0x2B, 0x22, 0x1D)
BODY = RGBColor(0x4A, 0x3B, 0x31)
CARAMEL = RGBColor(0xD0, 0x95, 0x5F)
CARAMEL_DEEP = RGBColor(0x9A, 0x5F, 0x33)
MUTED = RGBColor(0x8A, 0x72, 0x60)
MUTED_LIGHT = RGBColor(0xC9, 0xB2, 0x9B)

# Layout grid (inches, 26.67 x 15 canvas)
L, R = 1.8, 24.87
CW = R - L
GUT = 0.6
TOP = 3.9
BOT = 13.25
PAD = 0.55

prs = Presentation(TEMPLATE)
TEMPLATE_SLIDES = list(prs.slides)
BLANK = prs.slide_layouts[0]
page = 0
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


# ------------------------------------------------------------------ template cloning

def is_logo(sh):
    return sh.left is not None and sh.left >= Inches(25.5) and sh.top >= Inches(11.4)


def clone(index, keep=()):
    src = TEMPLATE_SLIDES[index - 1]
    dst = prs.slides.add_slide(BLANK)
    bg = src._element.cSld.bg
    if bg is not None:
        dst._element.cSld.insert(0, copy.deepcopy(bg))
    for sh in src.shapes:
        if sh.name not in keep or is_logo(sh):
            continue
        el = copy.deepcopy(sh._element)
        for node in el.iter():
            for attr, val in list(node.attrib.items()):
                if attr.startswith("{%s}" % R_NS):
                    rel = src.part.rels[val]
                    node.set(attr, dst.part.relate_to(rel.target_part, rel.reltype))
        dst.shapes._spTree.append(el)
    return dst


def move(slide, name, left=None, top=None):
    for sh in slide.shapes:
        if sh.name == name:
            if left is not None:
                sh.left = Inches(left)
            if top is not None:
                sh.top = Inches(top)


# ------------------------------------------------------------------ text

def _add_runs(p, txt, size, color, font, spc, italic=False):
    for part in re.split(r"(\*\*[^*]+\*\*)", txt):
        if not part:
            continue
        strong = part.startswith("**") and part.endswith("**")
        r = p.add_run()
        r.text = part[2:-2] if strong else part
        f = r.font
        f.size = Pt(size)
        f.color.rgb = color
        if italic:
            f.italic = True
        if strong and font == FONT:
            f.name, f.bold = FONT_SEMI, False
        else:
            f.name, f.bold = font, False
        if spc:
            f._rPr.set("spc", str(spc))


def _bullet(p, size, color, char="•"):
    pPr = p._p.get_or_add_pPr()
    indent = int(Pt(size) * 0.95)
    pPr.set("marL", str(indent))
    pPr.set("indent", str(-indent))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", str(color))
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def text(slide, x, y, w, h, paras, size=26, color=BODY, font=FONT, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, after=0, line=1.18, bullet=False, spc=0, name=None):
    """paras: str or list of str / (str, overrides). `**x**` = semibold run.
    Override keys: size, color, font, after, before, line, align, bullet, spc."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(paras, str):
        paras = [paras]
    for i, item in enumerate(paras):
        t, o = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = o.get("align", align)
        p.line_spacing = o.get("line", line)
        p.space_after = Pt(o.get("after", after))
        if o.get("before"):
            p.space_before = Pt(o["before"])
        sz = o.get("size", size)
        if o.get("bullet", bullet):
            _bullet(p, sz, CARAMEL)
        _add_runs(p, t, sz, o.get("color", color), o.get("font", font), o.get("spc", spc), o.get("italic", False))
    return tb


def bullets(slide, x, y, w, h, items, size=24, color=BODY, after=12, **kw):
    return text(slide, x, y, w, h, [(i, {"bullet": True}) for i in items], size=size, color=color,
                after=after, **kw)


# ------------------------------------------------------------------ shapes

def box(slide, x, y, w, h, fill=PAPER, line=None, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        dash=False, line_w=1.25):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = MSO_LINE.DASH
    s.shadow.inherit = False
    return s


def rule(slide, x, y, w, color=RULE, weight=1.25):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    c.line.color.rgb = color
    c.line.width = Pt(weight)
    return c


def arrow(slide, x1, y1, x2, y2, color=DARK, width=2.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    tail = etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


def photo(slide, path, x, y, w, h, focus=(0.5, 0.5), rounded=True, frame=None):
    """Crops the image to the box's aspect around `focus` and places it."""
    im = Image.open(path)
    target = w / h
    iw, ih = im.size
    if iw / ih > target:
        nw = int(ih * target)
        left = int((iw - nw) * focus[0])
        im = im.crop((left, 0, left + nw, ih))
    else:
        nh = int(iw / target)
        top = int((ih - nh) * focus[1])
        im = im.crop((0, top, iw, top + nh))
    out = os.path.join(TMP, f"crop_{os.path.basename(path)}_{w:.2f}x{h:.2f}.png")
    im.save(out)
    pic = slide.shapes.add_picture(out, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
        geom = pic._element.spPr.find(qn("a:prstGeom"))
        av = geom.find(qn("a:avLst"))
        if av is None:
            av = etree.SubElement(geom, qn("a:avLst"))
        gd = etree.SubElement(av, qn("a:gd"))
        gd.set("name", "adj")
        gd.set("fmla", "val %d" % int(0.035 * 100000 * max(w, h) / min(w, h)))
    if frame is not None:
        pic.line.color.rgb = frame
        pic.line.width = Pt(1.5)
    return pic


def placeholder(slide, x, y, w, h, what, dark=False):
    s = box(slide, x, y, w, h, fill=DARK_ALT if dark else PAPER, line=CARAMEL, dash=True, line_w=2.25,
            radius=0.04)
    s.name = "PLACEHOLDER - " + what[:40]
    text(slide, x + PAD, y, w - 2 * PAD, h, [
        ("ADD SCREENSHOT", {"size": 17, "font": FONT_SEMI, "color": CARAMEL, "spc": 200, "after": 10}),
        (what, {"size": 22, "color": CREAM if dark else BODY}),
    ], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line=1.2)


# ------------------------------------------------------------------ slide furniture

def header(slide, label, title, dark, size=58, x=L, w=CW):
    text(slide, x, 1.15, w, 0.45, label.upper(), size=17, font=FONT_SEMI, spc=250,
         color=CARAMEL if dark else CARAMEL_DEEP)
    text(slide, x, 1.6, w, 1.5, title, size=size, font=FONT_SEMI, color=CREAM if dark else INK, line=1.0)


def footer(slide, dark, x=L):
    global page
    page += 1
    col = MUTED_LIGHT if dark else MUTED
    text(slide, x, 14.0, 12, 0.4, "XR Continuum  ·  Barista Calibration Trainer", size=15, color=col)
    text(slide, R - 1.5, 14.0, 1.5, 0.4, f"{page:02d}", size=15, font=FONT_SEMI, color=col, align=PP_ALIGN.RIGHT)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def est_lines(t, size, w):
    t = t.replace("**", "")
    cpl = max(1, int(w / (size / 72 * 0.5)))
    lines, cur = 0, 0
    for word in t.split():
        if cur and cur + 1 + len(word) > cpl:
            lines += 1
            cur = len(word)
        else:
            cur += len(word) + (1 if cur else 0)
    return lines + 1


def est_h(items, size, w, line=1.22, after=10):
    items = [items] if isinstance(items, str) else items
    return sum(est_lines(t, size, w) * size / 72 * line * 1.2 for t in items) + (len(items) - 1) * after / 72


def card_h(w, head, body, kicker=None, head_size=28, body_size=23, as_bullets=False):
    inner = w - 2 * PAD - (body_size / 72 * 0.95 if as_bullets else 0)
    h = 2 * PAD + (0.5 if kicker else 0) + head_size / 72 * 1.05 + 0.35
    return h + est_h(body, body_size, inner, after=10 if as_bullets else 0)


def card(slide, x, y, w, h, head, body, dark, kicker=None, head_size=28, body_size=23, as_bullets=False):
    if h is None:
        h = card_h(w, head, body, kicker, head_size, body_size, as_bullets)
    box(slide, x, y, w, h, fill=DARK_CARD if dark else PAPER)
    cy = y + PAD
    if kicker:
        text(slide, x + PAD, cy, w - 2 * PAD, 0.4, kicker, size=16, font=FONT_SEMI, spc=200,
             color=CARAMEL if dark else CARAMEL_DEEP)
        cy += 0.5
    text(slide, x + PAD, cy, w - 2 * PAD, 0.6, head, size=head_size, font=FONT_SEMI,
         color=CREAM if dark else INK, line=1.0)
    cy += head_size / 72 * 1.05 + 0.35
    col = MUTED_LIGHT if dark else BODY
    if as_bullets:
        bullets(slide, x + PAD, cy, w - 2 * PAD, y + h - PAD - cy, body, size=body_size, color=col, after=10)
    else:
        text(slide, x + PAD, cy, w - 2 * PAD, y + h - PAD - cy, body, size=body_size, color=col, line=1.22)


def grid_x(cols, i, gut=GUT):
    w = (CW - gut * (cols - 1)) / cols
    return L + i * (w + gut), w


def table(slide, x, y, col_w, rows, row_h, dark, size=19, head_size=17):
    gf = slide.shapes.add_table(len(rows), len(col_w), Inches(x), Inches(y),
                                Inches(sum(col_w)), Inches(row_h * len(rows)))
    tblPr = gf._element.graphic.graphicData.tbl.tblPr
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    tbl = gf.table
    for i, wv in enumerate(col_w):
        tbl.columns[i].width = Inches(wv)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(row_h * (0.75 if r == 0 else 1))
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = DARK if not dark else CARAMEL_DEEP
            else:
                cell.fill.fore_color.rgb = (DARK_CARD if r % 2 else DARK_ALT) if dark else \
                                           (PAPER if r % 2 else PAPER_ALT)
            cell.margin_left = cell.margin_right = Inches(0.3)
            cell.margin_top = cell.margin_bottom = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.line_spacing = 1.12
            if r == 0:
                _add_runs(p, val.upper(), head_size, CREAM, FONT_SEMI, 150)
            else:
                _add_runs(p, val, size, CREAM if dark else BODY, FONT_SEMI if c == 0 else FONT, 0)
    return gf


def barista_png():
    im = Image.open(BARISTA)
    im = im.crop(im.split()[-1].getbbox())
    path = os.path.join(TMP, "barista_cropped.png")
    im.save(path)
    return path, im.size[0] / im.size[1]


BARISTA_IMG, BARISTA_ASPECT = barista_png()


def add_barista(slide, x, bottom, height):
    w = height * BARISTA_ASPECT
    pic = slide.shapes.add_picture(BARISTA_IMG, Inches(x), Inches(bottom - height), Inches(w), Inches(height))
    pic.name = "Barista illustration"


def shot(name):
    return os.path.join(SHOTS, name + ".png")


def spotlight_png():
    """Warm pool of light falling off to near-black, like a lone lamp in a dark pitch room."""
    import numpy as np
    w, h = 1600, 900
    yy, xx = np.mgrid[0:h, 0:w].astype(np.float32)
    d = np.sqrt(((xx - w * 0.5) / (w * 0.62)) ** 2 + ((yy - h * 0.46) / (h * 0.78)) ** 2)
    t = np.clip(d, 0, 1) ** 1.6
    inner, outer = np.array([0x4B, 0x30, 0x22]), np.array([0x12, 0x0B, 0x08])
    rgb = inner * (1 - t[..., None]) + outer * t[..., None]
    rgb += np.random.default_rng(7).normal(0, 2.2, rgb.shape)
    path = os.path.join(TMP, "spotlight.png")
    Image.fromarray(np.clip(rgb, 0, 255).astype("uint8")).save(path)
    return path


SPOTLIGHT = spotlight_png()


def fade(slide, speed="slow"):
    el = slide._element
    tr = etree.Element(qn("p:transition"))
    tr.set("spd", speed)
    etree.SubElement(tr, qn("p:fade"))
    anchor = el.find(qn("p:clrMapOvr"))
    if anchor is None:
        anchor = el.find(qn("p:cSld"))
    anchor.addnext(tr)


def dark_room():
    s = prs.slides.add_slide(BLANK)
    s.shapes.add_picture(SPOTLIGHT, 0, 0, prs.slide_width, prs.slide_height)
    fade(s)
    return s


def line(slide, y, h, t, size=84, color=CREAM, font=FONT_SEMI, italic=False):
    return text(slide, 2.5, y, 26.67 - 5.0, h, [(t, {"italic": italic})], size=size, color=color, font=font,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line=1.08)


# ================================================================== SLIDES

# 0 · Cold open (part of segment 1, about 45 seconds)
s = dark_room()
line(s, 4.6, 4.0, "Do you want to brew a coffee?", size=88)
text(s, 0, 12.6, 26.67, 0.5, "XR CONTINUUM PRESENTS", size=17, font=FONT_SEMI, spc=400, color=MUTED_LIGHT,
     align=PP_ALIGN.CENTER)
page += 1
notes(s, "COLD OPEN. Lights low if you can. One speaker, standing, no reading. Click, then wait a full two "
         "seconds before speaking.\n\n\"Do you want to brew a coffee?\"\n\n[Pause. Let them answer in their heads.]")

s = dark_room()
line(s, 3.3, 3.2, "A double shot flat white.", size=88)
specs = [("18 g", "coffee in"), ("25–30 s", "extraction"), ("65 °C", "milk")]
sw = 5.2
sx0 = (26.67 - 3 * sw) / 2
for i, (n, lab) in enumerate(specs):
    x = sx0 + i * sw
    if i:
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(8.3), Inches(x), Inches(10.4))
        c.line.color.rgb = DARK_RULE
        c.line.width = Pt(1.25)
    text(s, x, 8.1, sw, 2.5, [(n, {"size": 54, "font": FONT_SEMI, "color": CARAMEL, "after": 4}),
                              (lab.upper(), {"size": 15, "font": FONT_SEMI, "spc": 300, "color": MUTED_LIGHT})],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line=1.05)
page += 1
notes(s, "\"A double shot flat white.\"\n\n[Slower now, almost fond.] \"Eighteen grams of coffee. Twenty-five to "
         "thirty seconds. Milk at sixty-five degrees.\"\n\n[It should sound simple. That's the trap.]")

s = dark_room()
line(s, 5.0, 3.6, "But do you know how?", size=96, color=CARAMEL, font=FONT, italic=True)
page += 1
notes(s, "[Beat. Then quietly:] \"But do you know how?\"\n\n[Hold eye contact with the panel. Count three.]")

s = dark_room()
line(s, 3.9, 2.6, "Nobody does, the first time.", size=80)
line(s, 7.0, 2.6, "Every barista learned on a live machine, in the middle of a rush, "
                  "one wasted shot at a time.", size=34, color=MUTED_LIGHT, font=FONT)
page += 1
notes(s, "\"Nobody does, the first time. Every barista you've ever met learned on a live machine, in the middle "
         "of a rush, one wasted shot at a time. Beans in the bin. A queue at the counter. Someone senior sighing "
         "over their shoulder.\"")

s = dark_room()
line(s, 3.9, 2.4, "This isn't a game.", size=80)
line(s, 6.4, 2.6, "It's the morning before your first shift.", size=80, color=CARAMEL)
page += 1
notes(s, "\"So we built something. This isn't a game. It's the morning before your first shift: the café is "
         "empty, the machine is warm, and every mistake is free.\"\n\n[Click to the title. Lights up.]")

# 1 · Title
s = clone(1, {"Image 0", "Image 1", "Image 3", "Image 5"})
fade(s)
move(s, "Image 5", left=0.3, top=10.6)
add_barista(s, 17.2, 15.0, 13.6)
text(s, L, 2.75, 15, 0.45, "INTE 42312  ·  VIRTUAL & AUGMENTED REALITY", size=18, font=FONT_SEMI, spc=250,
     color=CARAMEL_DEEP)
text(s, L, 3.4, 15.2, 4.4, "Barista Calibration Trainer", size=108, font=FONT_SEMI, color=INK, line=0.92)
text(s, L, 7.85, 14.5, 0.8, "Dialling in a commercial espresso machine in VR", size=34, color=BODY)
rule(s, L, 9.15, 1.4, color=CARAMEL, weight=3)
text(s, L, 9.5, 14.5, 0.45, "TEAM XR CONTINUUM", size=18, font=FONT_SEMI, spc=250, color=CARAMEL_DEEP)
text(s, L, 10.0, 14.5, 0.6, "<Member 1>   ·   <Member 2>   ·   <Member 3>   ·   <Member 4>", size=26, color=INK)
page += 1
notes(s, "Segment 1, lights up. \"This is the Barista Calibration Trainer, by XR Continuum, for INTE 42312.\" "
         "Introduce every member by name, then go straight to the contribution slide: the brief requires the "
         "declaration and timeline first, so the cold open must stay under a minute.")

# 2 · Individual contribution
s = clone(4, {"Image 0"})
header(s, "02 · Individual contribution", "Who built what, and who defends it", True)
members = [
    ("MEMBER A", "<name>", "XR rig & build",
     ["Device tracking origin, seated eye height", "Snap turn, no locomotion", "Keyboard & mouse fallback",
      "Windows build pipeline"], "CafeSceneBuilder (rig), DesktopRigDriver, input actions"),
    ("MEMBER B", "<name>", "Dose & tamp",
     ["Grinder dosing & grind dial", "Upright tamper grip", "Spring tamp: force and level", "Portafilter sockets"],
     "GrinderStation, GrindDial, Tamper, TamperPress, TamperGrabTransformer"),
    ("MEMBER C", "<name>", "Brew & model",
     ["Group-head bayonet lock", "Live extraction model", "Pressure gauge, timer, yield"],
     "GroupHead, ExtractionModel, MachineGauge"),
    ("MEMBER D", "<name>", "Milk, flow & audio",
     ["Steam to 65 °C and pour", "ShiftFlow, ticket, result dial", "Spatial audio", "Welcome & credits boards"],
     "SteamWand, MilkPour, ShiftFlow, OrderTicket, ResultDial, AudioKit"),
]
_, mw = grid_x(4, 0, 0.45)
b_h = max(est_h(m[3], 24, mw - 2 * PAD - 0.3) for m in members)
f_h = max(est_h(m[4], 18, mw - 2 * PAD) for m in members)
mh = PAD + 2.55 + b_h + 0.55 + 0.45 + f_h + PAD
for i, (tag, name, area, items, files) in enumerate(members):
    x, w = grid_x(4, i, 0.45)
    y = TOP
    box(s, x, y, w, mh, fill=DARK_CARD)
    text(s, x + PAD, y + PAD, w - 2 * PAD, 0.4, tag, size=16, font=FONT_SEMI, spc=200, color=CARAMEL)
    text(s, x + PAD, y + PAD + 0.5, w - 2 * PAD, 0.5, name, size=22, color=MUTED_LIGHT)
    text(s, x + PAD, y + PAD + 1.15, w - 2 * PAD, 0.6, area, size=32, font=FONT_SEMI, color=CREAM)
    rule(s, x + PAD, y + PAD + 2.2, w - 2 * PAD, color=DARK_RULE)
    bullets(s, x + PAD, y + PAD + 2.55, w - 2 * PAD, b_h, items, size=24, color=CREAM, after=10)
    fy = y + PAD + 2.55 + b_h + 0.55
    text(s, x + PAD, fy, w - 2 * PAD, 0.4, "KEY FILES", size=14, font=FONT_SEMI, spc=200, color=CARAMEL)
    text(s, x + PAD, fy + 0.45, w - 2 * PAD, f_h, files, size=18, color=MUTED_LIGHT, line=1.2)
text(s, L, TOP + mh + 0.55, CW, 0.5, "Each member operates and explains their own part live in the demo.",
     size=21, color=MUTED_LIGHT)
footer(s, True)
notes(s, "Segment 2, presented FIRST. Each person names their part and files. The panel may question anyone on "
         "their own part, so only claim what you can explain. With three members, merge B and C.")

# 3 · Timeline
s = clone(14, {"Image 0"})
header(s, "03 · Project management", "Plan vs. what actually happened", False)
lx = R - 9.2
box(s, lx, 2.3, 0.75, 0.28, fill=DARK, radius=0.5)
text(s, lx + 0.95, 2.18, 2.6, 0.5, "Planned", size=19, color=BODY)
box(s, lx + 3.4, 2.3, 0.75, 0.28, fill=CARAMEL, radius=0.5)
text(s, lx + 4.35, 2.18, 4.9, 0.5, "Actual (drag to match)", size=19, color=BODY)
days = ["Mon 5", "Tue 6", "Wed 7", "Thu 8", "Fri 9", "Sat 10"]
tasks = [("Rig, packages, grey-box café", 0, 1), ("Dose & tamp", 1, 1), ("Lock, extraction model, gauge", 2, 1),
         ("Steam, pour, flow, audio", 3, 1), ("UX, physics & text polish", 3, 2),
         ("Windows build, cross-testing, credits", 4, 1), ("Rehearse, submit & present", 5, 1)]
lab_w = 7.4
gx = L + lab_w
gw = (R - gx) / 6
hy = TOP - 0.1
text(s, L, hy, lab_w, 0.5, "TASK", size=15, font=FONT_SEMI, spc=200, color=MUTED)
for d, label in enumerate(days):
    text(s, gx + d * gw, hy, gw, 0.5, (label + " OCT").upper(), size=15, font=FONT_SEMI, spc=150, color=MUTED,
         align=PP_ALIGN.CENTER)
rule(s, L, hy + 0.6, CW)
rh = 1.08
for r, (name, st, ln) in enumerate(tasks):
    y = hy + 0.75 + r * rh
    if r % 2 == 0:
        box(s, L - 0.2, y, CW + 0.4, rh, fill=PAPER, radius=0.12)
    text(s, L, y, lab_w - 0.3, rh, name, size=21, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    box(s, gx + st * gw + 0.15, y + 0.24, ln * gw - 0.3, 0.26, fill=DARK, radius=0.5)
    b = box(s, gx + st * gw + 0.15, y + 0.6, ln * gw - 0.3, 0.26, fill=CARAMEL, radius=0.5)
    b.name = "ACTUAL bar - " + name
dy = hy + 0.75 + len(tasks) * rh + 0.35
box(s, L, dy, 0.09, BOT - dy, fill=CARAMEL, shape=MSO_SHAPE.RECTANGLE)
text(s, L + 0.4, dy, CW - 0.4, BOT - dy, [
    ("DEVIATIONS & WHY", {"size": 15, "font": FONT_SEMI, "spc": 200, "color": CARAMEL_DEEP, "after": 6}),
    ("<e.g. Unity 6 + XR packages took most of day 1 · text and physics feel needed an extra polish pass after "
     "the first simulator tests · the exe showed an old scene until the build menu regenerated it>",
     {"size": 21, "color": BODY}),
], anchor=MSO_ANCHOR.MIDDLE, line=1.2)
footer(s, False)
notes(s, "Segment 3, straight after the declaration. Dark bars = plan, caramel = actual. Drag the caramel bars "
         "to reality and give one reason per slip; honesty about deviations is what is marked.")

# 4 · Problem
s = clone(13, {"Image 0"})
header(s, "04 · Problem statement", "Baristas learn on the live machine", True)
text(s, L, TOP + 0.1, 10.6, BOT - TOP, [
    ("New café staff learn espresso on the café's only machine, often during service.", {"after": 22}),
    ("Every practice shot uses about 18 g of specialty coffee and every practice latte uses milk. "
     "Dialling in a grinder takes many shots.", {"after": 22}),
    ("The group head is blocked, customers wait, and a shift lead has to stand beside the trainee.", {}),
], size=28, color=CREAM, line=1.25)
stats = [("18 g", "of coffee binned on every practice shot"),
         ("1", "group head tied up while a trainee practises"),
         ("3", "people affected: trainee, shift lead, owner")]
sx = 14.3
for i, (num, lab) in enumerate(stats):
    y = TOP + i * 3.1
    if i:
        rule(s, sx, y, R - sx, color=DARK_RULE)
    text(s, sx, y + 0.3, 4.3, 2.4, num, size=100, font=FONT_SEMI, color=CARAMEL, anchor=MSO_ANCHOR.MIDDLE)
    text(s, sx + 4.6, y + 0.3, R - sx - 4.6, 2.4, lab, size=26, color=CREAM, anchor=MSO_ANCHOR.MIDDLE, line=1.2)
footer(s, True)
notes(s, "Segment 4: the specific problem, who it affects and why it matters. Keep it concrete: beans, milk, a "
         "blocked machine, a supervisor's time.")

# 5 · Vision
s = clone(7, {"Image 0", "Image 2"})
text(s, 6.4, 1.15, 16, 0.45, "05 · VISION STATEMENT", size=17, font=FONT_SEMI, spc=250, color=CARAMEL)
text(s, 6.4, 3.7, 16.4, 6.0,
     "A new barista sits down, reads one order ticket and pulls a correctly calibrated flat white, without "
     "anyone showing them how and without wasting a single real bean.",
     size=50, font=FONT_SEMI, color=CREAM, line=1.15, anchor=MSO_ANCHOR.MIDDLE)
rule(s, 6.4, 10.35, 1.4, color=CARAMEL, weight=3)
text(s, 6.4, 10.65, 16, 0.6, "What success looks like for the trainee", size=26, color=MUTED_LIGHT)
footer(s, True)
notes(s, "Segment 5: read it once, slowly. Every later slide traces back to: unaided, calibrated, no waste.")

# 6 · Objectives
s = clone(17, {"Image 0"})
box(s, 0, 0, 8.4, 15.0, fill=DARK, shape=MSO_SHAPE.RECTANGLE)
add_barista(s, 0.85, 15.0, 12.6)
ox = 9.7
header(s, "06 · Objectives", "Three measurable objectives", False, x=ox, w=R - ox)
objs = [("Unaided, end to end",
         "A first-time user completes the order, ticket to verdict, with no help in **under 6 minutes**."),
        ("Understands calibration",
         "Explains a sour or bitter shot from the result dial and reaches **Balanced within 3 attempts**."),
        ("Milk on temperature", "Steamed milk finishes between **63 and 67 °C** on at least one attempt.")]
oh = (BOT - TOP - 2 * 0.4) / 3
for i, (head, body) in enumerate(objs):
    y = TOP + i * (oh + 0.4)
    box(s, ox, y, R - ox, oh, fill=PAPER)
    text(s, ox + PAD, y, 1.6, oh, f"0{i + 1}", size=60, font=FONT_SEMI, color=CARAMEL, anchor=MSO_ANCHOR.MIDDLE)
    text(s, ox + 2.5, y, R - ox - 2.5 - PAD, oh, [
        (head, {"size": 30, "font": FONT_SEMI, "color": INK, "after": 8}),
        (body, {"size": 23, "color": BODY}),
    ], anchor=MSO_ANCHOR.MIDDLE, line=1.2)
footer(s, False, x=ox)
notes(s, "Segment 6: up to three, each measurable. Results come back on the conclusion slide.")

# 7 · Theory
s = clone(12, {"Image 0"})
header(s, "07 · Theoretical grounding", "Module theory, applied to real decisions", False)
theory = [
    ("Presence", "One coherent café, no menus floating in a void. Instructions on a paper ticket, the gauge on "
                 "the machine, sounds from their source."),
    ("Cybersickness", "Seated and stationary, so no vection. The only rotation is an instant 45° snap turn; "
                      "smooth move and teleport are off."),
    ("Vergence-accommodation", "Reading panels sit 0.75–1 m away. Tools are within arm's reach but only "
                               "glanced at, not read."),
    ("D.I.C.E.", "Practising on the real machine is Counterproductive (blocks service) and Expensive (beans "
                 "and milk)."),
    ("Embodied learning", "Dose, grind, tamp force and level come from the hands, not a menu, so practice "
                          "transfers to the real task."),
    ("Spatial audio", "Every sound is a 3D source at its object: VR's default, the reverse of the mono AR "
                      "convention."),
]
_, tw3 = grid_x(3, 0)
ch = min((BOT - TOP - GUT) / 2, max(card_h(tw3, h, b, "01", 30, 25) for h, b in theory))
for i, (h, b) in enumerate(theory):
    x, w = grid_x(3, i % 3)
    card(s, x, TOP + (i // 3) * (ch + GUT), w, ch, h, b, False, kicker=f"0{i + 1}", head_size=30, body_size=25)
footer(s, False)
notes(s, "Segment 7: for each concept say the decision it caused. The criterion is application, not recital.")

# 8 · Concept overview
s = clone(15, {"Image 0", "Image 2"})
header(s, "08 · Proposed solution", "One shift, one café, one flat white", False)
photo(s, shot("cafe_overview"), 2.39, 3.93, 10.82, 7.10, rounded=False)
tx = 15.6
steps8 = [("The brief", "You're the new barista. An order ticket asks for a flat white."),
          ("The work", "Unlock, dose, tamp, lock in, pull the shot, steam milk to 65 °C, pour."),
          ("The verdict", "A result dial grades the shot Sour, Balanced, Bitter or Channeled and says what "
                          "to change."),
          ("The twist", "The grinder starts coarse on purpose, so shot one runs sour. You calibrate, like a "
                        "real barista every morning.")]
for i, (h, b) in enumerate(steps8):
    y = TOP + i * 2.35
    text(s, tx, y, R - tx, 2.2, [(h.upper(), {"size": 15, "font": FONT_SEMI, "spc": 200, "color": CARAMEL_DEEP,
                                              "after": 6}),
                                 (b, {"size": 25, "color": INK})], line=1.2)
text(s, 2.39, 12.75, 10.82, 0.5, "Seated start in the exported build: welcome board, station panel, menu, machine.",
     size=17, color=MUTED, align=PP_ALIGN.CENTER)
footer(s, False)
notes(s, "Segment 8: one story, not a feature list. Emphasise the calibration loop: first shot sour by design.")

# 9 · Shift flow
s = clone(14, {"Image 0"})
header(s, "08 · Concept: the shift", "A clear beginning, middle and end", False)
steps = [("Briefing", "Welcome board, START"), ("Detach", "Free the portafilter"), ("Dose", "Hold to grind, ~18 g"),
         ("Tamp", "Level press, force shown"), ("Lock", "Bayonet into group head"), ("Brew", "Live gauge, time, yield"),
         ("Steam", "Stop at 65 °C"), ("Pour", "Tilt pitcher over cup"), ("Result", "Verdict + what to change")]
nw, nh, ng = 4.0, 1.9, (CW - 5 * 4.0) / 4
cols = [L + i * (nw + ng) for i in range(5)]
r1, r2 = TOP, TOP + nh + 1.5
pos = {i: (cols[i], r1) for i in range(5)}
for j, i in enumerate(range(5, 9)):
    pos[i] = (cols[4 - j], r2)
for i, (h, d) in enumerate(steps):
    x, y = pos[i]
    end = i in (0, 8)
    box(s, x, y, nw, nh, fill=DARK if end else PAPER, radius=0.14)
    text(s, x + 0.3, y, nw - 0.6, nh, [
        (h, {"size": 28, "font": FONT_SEMI, "color": CARAMEL if end else INK, "after": 4}),
        (d, {"size": 18, "color": MUTED_LIGHT if end else BODY}),
    ], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line=1.15)
for i in range(4):
    arrow(s, cols[i] + nw + 0.1, r1 + nh / 2, cols[i + 1] - 0.1, r1 + nh / 2)
    arrow(s, cols[4 - i] - 0.1, r2 + nh / 2, cols[3 - i] + nw + 0.1, r2 + nh / 2)
arrow(s, cols[4] + nw / 2, r1 + nh + 0.1, cols[4] + nw / 2, r2 - 0.1)
arrow(s, cols[1] + nw / 2, r2 - 0.1, cols[1] + nw / 2, r1 + nh + 0.1, color=CARAMEL_DEEP)
text(s, cols[1] + nw / 2 + 0.25, r1 + nh + 0.2, 4.5, 1.1, [
    ("Pull another shot", {"font": FONT_SEMI, "color": CARAMEL_DEEP}), ("grind setting is kept", {"color": MUTED})],
     size=18, anchor=MSO_ANCHOR.MIDDLE, line=1.1)
cy = r2 + nh + 0.75
for i, (h, b) in enumerate([
        ("Clear start", "The welcome board and order ticket explain the goal in seconds; a glowing marker "
                        "points at the next thing to touch."),
        ("Real conclusion", "The result dial resolves the order and says how to improve. Reset, Exit and "
                            "Credits stay on the station panel.")]):
    x, w = grid_x(2, i)
    card(s, x, cy, w, BOT - cy, h, b, False, head_size=26, body_size=22)
footer(s, False)
notes(s, "Still segment 8. Walk the loop: clear start, clear end, retry keeps the grind setting.")

# 10 · Bridge decision 1
s = clone(4, {"Image 0", "Image 1"})
header(s, "09 · Bridge decision 1", "Tracking origin: Device, seated", True)
lw = 10.6
bullets(s, L, TOP, lw, 4.6, [
    "The task is a fixed reach envelope in front of a seated person.",
    "Device origin places the counter relative to where the headset starts.",
    "A 1.2 m Camera Y Offset gives the same seated eye height in the simulator and on a headset.",
], size=24, color=CREAM, after=14, line=1.2)
iw, ih = 8.7, 8.7 * (198 / 456)
photo(s, shot("xr_origin"), L, 8.65, iw, ih, frame=DARK_RULE)
text(s, L, 8.65 + ih + 0.12, lw, 0.4,
     "XR Origin on the rig: Tracking Origin Mode Device, Camera Y Offset 1.2",
     size=16, color=MUTED_LIGHT)
rx = L + lw + GUT + 0.3
rw = R - rx
ch1 = 3.6
box(s, rx, TOP, rw, ch1, fill=DARK_CARD)
text(s, rx + PAD, TOP + PAD, rw - 2 * PAD, ch1 - 2 * PAD, [
    ("NOT CHOSEN", {"size": 15, "font": FONT_SEMI, "spc": 200, "color": MUTED_LIGHT, "after": 6}),
    ("Floor-referenced", {"size": 30, "font": FONT_SEMI, "color": CREAM, "after": 10}),
    ("Virtual floor = real floor, ideal for standing and room-scale. A seated user would end up with their "
     "eyes at counter height.", {"size": 22, "color": MUTED_LIGHT}),
], line=1.2)
y2 = TOP + ch1 + 0.45
box(s, rx, y2, rw, BOT - y2, fill=CARAMEL)
text(s, rx + PAD, y2 + PAD, rw - 2 * PAD, BOT - y2 - 2 * PAD, [
    ("OUR CHOICE", {"size": 15, "font": FONT_SEMI, "spc": 200, "color": DARK, "after": 6}),
    ("Device-referenced", {"size": 30, "font": FONT_SEMI, "color": DARK, "after": 12}),
    ("Origin = where the headset started, ideal for a seated workstation.", {"size": 22, "color": DARK,
                                                                             "after": 10}),
    ("Consistent 1.2 m eye height over a 0.84 m counter.", {"size": 22, "color": DARK, "bullet": True,
                                                            "after": 6}),
    ("Works at a desk with the simulator and in a seated headset.", {"size": 22, "color": DARK,
                                                                     "bullet": True}),
], line=1.2)
footer(s, True)
notes(s, "Bridge decision 1. Likely question: what if the user stands? The counter stays relative to the start "
         "pose; the experience is designed and labelled as seated.")

# 11 · Bridge decision 2
s = clone(12, {"Image 0"})
header(s, "09 · Bridge decision 2", "Why VR, not AR: D.I.C.E.", False)
dice = [("D", "Dangerous", False, "Hot steam and a 9-bar group head, but a minor factor here."),
        ("I", "Impossible", False, "A real machine can't show live pressure or extraction. VR can."),
        ("C", "Counterproductive", True, "Practice blocks the only group head and delays customers."),
        ("E", "Expensive", True, "About 18 g of beans, plus milk, wasted on every practice shot.")]
th = 5.6
for i, (letter, word, main, why) in enumerate(dice):
    x, w = grid_x(4, i, 0.45)
    box(s, x, TOP, w, th, fill=DARK if main else PAPER)
    text(s, x + PAD, TOP + 0.25, 2.2, 2.0, letter, size=110, font=FONT_SEMI, color=CARAMEL)
    text(s, x + w - PAD - 3.0, TOP + PAD, 3.0, 0.4, "MAIN REASON" if main else "PARTLY", size=14,
         font=FONT_SEMI, spc=200, color=CARAMEL if main else MUTED, align=PP_ALIGN.RIGHT)
    text(s, x + PAD, TOP + 2.45, w - 2 * PAD, th - 2.45 - PAD, [
        (word, {"size": 27, "font": FONT_SEMI, "color": CREAM if main else INK, "after": 8}),
        (why, {"size": 21, "color": MUTED_LIGHT if main else BODY}),
    ], line=1.2)
by = TOP + th + 0.5
box(s, L, by, CW, BOT - by, fill=PAPER)
text(s, L + PAD, by + PAD, 7.0, BOT - by - 2 * PAD,
     "What replacing the world gives us that adding to it cannot", size=27, font=FONT_SEMI, color=INK, line=1.15)
bullets(s, L + 8.3, by + PAD, CW - 8.3 - PAD, BOT - by - 2 * PAD, [
    "AR still needs the real machine, beans, milk and service floor. VR removes all four.",
    "We can show hidden state (live pressure, extraction time) and repeat a shot instantly.",
    "We keep the café's sound but remove its rush, so the trainee can focus.",
], size=22, color=BODY, after=8, line=1.18)
footer(s, False)
notes(s, "Bridge decision 2. Lead with C and E; admit D and I only partly apply. That honesty earns rationale marks.")

# 12 · Design rationale
s = clone(2, {"Image 0", "Image 1", "Image 3"})
header(s, "09 · Design rationale", "Comfort, interaction, UI and audio", True)
grid = [
    ("Locomotion & comfort", ["No movement provider: seated, everything within reach",
                              "45° snap turn only: instant, so no vection",
                              "No smooth move or teleport: the task never needs travel"]),
    ("Interaction", ["XRI grab interactables with velocity tracking, so tools feel weighty",
                     "Sockets hold the portafilter (fork, mat) and the pitcher (steam rest)",
                     "Dials and lock are hand-position drags; all input via named XRI actions"]),
    ("Spatial UI", ["Ticket, gauge, thermometer and result dial are objects in the café: no HUD",
                    "Reading panels at 0.75–1 m to limit vergence-accommodation strain",
                    "A glowing marker hovers over the next step"]),
    ("Audio", ["Every sound is a 3D source at its object, logarithmic roll-off",
               "Grinder, pump, steam hiss (pitch falls as milk heats), pour, knock, chimes",
               "All synthesised in code: nothing to license"]),
]
_, gw2 = grid_x(2, 0, 0.45)
ch = min((BOT - TOP - 0.45) / 2, max(card_h(gw2, h, b, None, 30, 25, True) for h, b in grid))
for i, (h, b) in enumerate(grid):
    x, w = grid_x(2, i % 2, 0.45)
    card(s, x, TOP + (i // 2) * (ch + 0.45), w, ch, h, b, True, head_size=30, body_size=25, as_bullets=True)
footer(s, True)
notes(s, "Segment 9: each decision with its reason. Name the comfort defaults: seated, snap turn, no artificial motion.")

# 13 · Architecture
s = clone(14, {"Image 0"})
header(s, "10 · System architecture", "From XR Origin to game logic", False)
c1, c1w = L, 6.7
c3w = 7.4
c3 = R - c3w
c2w = 6.6
c2 = c1 + c1w + (c3 - c1 - c1w - c2w) / 2


def node(x, y, w, h, head, sub="", style="card"):
    fill, hc, sc = {"card": (PAPER, INK, BODY), "dark": (DARK, CARAMEL, MUTED_LIGHT),
                    "accent": (CARAMEL, DARK, DARK), "inner": (DARK_CARD, CREAM, MUTED_LIGHT)}[style]
    box(s, x, y, w, h, fill=fill, radius=0.1)
    paras = [(head, {"size": 21, "font": FONT_SEMI, "color": hc, "after": 3})]
    if sub:
        paras.append((sub, {"size": 16, "color": sc}))
    text(s, x + 0.3, y, w - 0.6, h, paras, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line=1.12)


rig_h = 6.85
box(s, c1, TOP, c1w, rig_h, fill=DARK, radius=0.05)
text(s, c1 + 0.3, TOP + 0.35, c1w - 0.6, 0.4, "XR ORIGIN · DEVICE, SEATED", size=14, font=FONT_SEMI, spc=200,
     color=CARAMEL, align=PP_ALIGN.CENTER)
ih = 1.75
for k, (h, sub) in enumerate([("Main Camera", "Tracked Pose Driver"),
                              ("Left / Right controller", "Near-Far Interactor"),
                              ("Snap Turn Provider", "45°, the only locomotion")]):
    node(c1 + 0.35, TOP + 0.95 + k * (ih + 0.2), c1w - 0.7, ih, h, sub, "inner")
node(c1, TOP + rig_h + 0.45, c1w, 0.95, "Named XRI input actions", "", "accent")
node(c1, TOP + rig_h + 1.6, c1w, BOT - (TOP + rig_h + 1.6), "DesktopRigDriver", "adds keyboard & mouse bindings")
arrow(s, c1 + c1w / 2, TOP + rig_h + 0.4, c1 + c1w / 2, TOP + rig_h + 0.05)

col2 = [("XR Grab Interactable", "portafilter, tamper, pitcher, cup"),
        ("XR Simple Interactable", "buttons, dials, portafilter lock"),
        ("ToolSocket", "grinder fork, tamp mat, steam rest"),
        ("Physics helpers", "ToolSettle, TamperGrabTransformer")]
h2 = (BOT - TOP - 3 * 0.35) / 4
for k, (h, sub) in enumerate(col2):
    node(c2, TOP + k * (h2 + 0.35), c2w, h2, h, sub)
arrow(s, c1 + c1w + 0.1, TOP + 3.9, c2 - 0.1, TOP + 3.9)

col3 = [("Stations", "Grinder, TamperPress, GroupHead, SteamWand, MilkPour", "card", 2.0),
        ("ShiftFlow", "step machine: Briefing → Result", "dark", 1.6),
        ("ExtractionModel", "advanced feature", "accent", 1.6),
        ("World-space outputs", "OrderTicket, HintMarker, MachineGauge, ResultDial, AudioKit", "card", 2.0)]
gap3 = (BOT - TOP - sum(c[3] for c in col3)) / 3
y = TOP
ys = []
for h, sub, st, hh in col3:
    node(c3, y, c3w, hh, h, sub, st)
    ys.append((y, hh))
    y += hh + gap3
arrow(s, c2 + c2w + 0.1, TOP + 1.0, c3 - 0.1, TOP + 1.0)
for k in range(3):
    y0, h0 = ys[k]
    arrow(s, c3 + c3w / 2, y0 + h0 + 0.08, c3 + c3w / 2, ys[k + 1][0] - 0.08)
text(s, c3 + c3w / 2 + 0.25, ys[0][0] + ys[0][1] + 0.08, 3.4, gap3 - 0.16, "Complete(step)", size=15,
     color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
footer(s, False)
notes(s, "Segment 10. Hands are interactors, tools are interactables, sockets are interactors that hold tools. "
         "Stations report to ShiftFlow; ExtractionModel drives the gauge and verdict. The whole scene is generated "
         "by Barista > Build Cafe Scene (CafeSceneBuilder.cs).")

# 14 · Advanced feature
s = clone(13, {"Image 0"})
header(s, "10 · Advanced feature", "Live extraction model", True)
iw_, ih_ = 6.75, BOT - TOP - 0.55
ix = R - iw_
lw = ix - GUT - 0.3 - L
fh = 3.5
box(s, L, TOP, lw, fh, fill=DARK_CARD)
text(s, L + PAD, TOP + PAD, lw - 2 * PAD, fh - 2 * PAD, [
    ("SHOT TIME", {"size": 15, "font": FONT_SEMI, "spc": 200, "color": CARAMEL, "after": 6}),
    ("t = 27.5 s · e^(−0.22 (grind − 5)) · (1 + 0.015 (tamp − 15)) · (dose ⁄ 18)^1.6",
     {"size": 23, "font": FONT_SEMI, "color": CREAM, "after": 4}),
    ("× 0.8 if tamp < 5 kg   ·   × 0.75 if tilt > 5° (channeling)", {"size": 19, "color": MUTED_LIGHT, "after": 18}),
    ("PRESSURE", {"size": 15, "font": FONT_SEMI, "spc": 200, "color": CARAMEL, "after": 6}),
    ("p = 9 bar · √(t ⁄ 27.5), clamped to 2.5–11 bar", {"size": 23, "font": FONT_SEMI, "color": CREAM}),
], line=1.15)
cy = TOP + fh + 0.4
chips = [("Sour", "< 25 s"), ("Balanced", "25–30 s"), ("Bitter", "> 30 s"), ("Channeled", "tilt > 5°")]
cwid = (lw - 3 * 0.3) / 4
for i, (v, r_) in enumerate(chips):
    x = L + i * (cwid + 0.3)
    good = v == "Balanced"
    box(s, x, cy, cwid, 1.3, fill=CARAMEL if good else DARK_CARD, radius=0.25)
    text(s, x, cy, cwid, 1.3, [(v, {"size": 22, "font": FONT_SEMI, "color": DARK if good else CREAM}),
                               (r_, {"size": 17, "color": DARK if good else MUTED_LIGHT})],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line=1.1)
bullets(s, L, cy + 1.8, lw, BOT - cy - 1.8, [
    "Turns what the hands did into a shot the trainee can read.",
    "Drives the live gauge, timer and yield, then the verdict dial.",
    "Deterministic, so fully testable in the XR Device Simulator.",
    "A teaching model with the right direction of effect, not fluid physics.",
], size=22, color=CREAM, after=8)
photo(s, shot("result_report"), ix, TOP, iw_, ih_)
text(s, ix, TOP + ih_ + 0.15, iw_, 0.45, "Shot 1, recorded run: 19.1 s at grind 6.5, so Sour", size=16,
     color=MUTED_LIGHT, align=PP_ALIGN.CENTER)
footer(s, True)
notes(s, "Advanced feature. It closes the loop between hand technique and outcome. The recorded run matches the "
         "model: starting grind 6.5 gave 19.1 s, so Sour, and the board says grind finer. Be explicit that it is a "
         "teaching model, not physics.")

# 15 · Handling & physics
s = clone(12, {"Image 0"})
header(s, "10 · Handling & physics", "Weight without chaos", False)
phys = [("Plausible weight", ["Velocity-tracked grabs: heavier tools lag the hand",
                              "Velocity changes capped, smoothing on",
                              "Physics at 90 Hz, solver iterations 12 / 4"]),
        ("Nothing clips through", ["Grabbed tools stay physical, so the counter and machine block them",
                                   "Push-out capped at 0.25 m/s, so contact never flings a tool",
                                   "Collisions ignored only where parts must nest"]),
        ("Settling & the tamper", ["Released tools rest upright, or return home",
                                   "Tamper held vertical: 30% wrist follow, max 12°",
                                   "Tamp resistance is visual, a force readout and haptics"])]
_, pw3 = grid_x(3, 0)
ph = max(card_h(pw3, h, b, "01", 30, 25, True) for h, b in phys)
for i, (h, b) in enumerate(phys):
    x, w = grid_x(3, i)
    card(s, x, TOP, w, ph, h, b, False, kicker=f"0{i + 1}", head_size=30, body_size=25, as_bullets=True)
ky = TOP + ph + 0.6
box(s, L, ky, CW, BOT - ky, fill=DARK)
kstats = [("90 Hz", "physics step"), ("12 / 4", "solver iterations"), ("0.25 m/s", "push-out cap"),
          ("12°", "max tamper wrist tilt")]
kw = CW / 4
for i, (n, lab) in enumerate(kstats):
    x = L + i * kw
    if i:
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(ky + 0.45), Inches(x),
                                                Inches(BOT - 0.45))
        c.line.color.rgb = DARK_RULE
        c.line.width = Pt(1.25)
    text(s, x, ky, kw, BOT - ky, [(n, {"size": 44, "font": FONT_SEMI, "color": CARAMEL, "after": 2}),
                                  (lab, {"size": 19, "color": MUTED_LIGHT})],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line=1.05)
footer(s, False)
notes(s, "Maps to the 13-mark technical criterion. Be honest: controllers can't push back, so resistance is "
         "simulated with visuals, a force number and haptics.")

# 16 · AI disclosure
s = clone(4, {"Image 0"})
header(s, "10 · AI-tool disclosure", "How we used AI, and how much", True)
table(s, L, TOP, [4.4, 10.27, 4.2, 4.2], [
    ["Tool", "Used for", "How much", "Defended by"],
    ["Cursor (AI agent)", "Drafting the scene builder and C# scripts from our design; debugging physics, text "
                          "shader and build; first draft of slides and viva notes", "<High: ~X% of code>",
     "Each member, own part"],
    ["<Other, if any>", "<e.g. explaining XRI concepts, checking references>", "<Low / Medium>", "<name>"],
    ["Team, no AI", "Concept, D.I.C.E. case, objectives, design choices, testing, tuning, verifying each fix",
     "n/a", "All"],
], row_h=1.55, dark=True, size=19)
ry = TOP + 1.55 * 0.75 + 3 * 1.55 + 0.55
box(s, L, ry, CW, BOT - ry, fill=CARAMEL)
text(s, L + PAD, ry, 5.0, BOT - ry, "Responsible use", size=28, font=FONT_SEMI, color=DARK,
     anchor=MSO_ANCHOR.MIDDLE)
text(s, L + 6.0, ry + PAD, CW - 6.0 - PAD, BOT - ry - 2 * PAD,
     "We specified what to build, reviewed every change, tested it in the simulator and the exported exe, and "
     "rejected fixes that didn't work, such as kinematic grabs that let tools pass through. Every member can "
     "explain their part live.", size=21, color=DARK, anchor=MSO_ANCHOR.MIDDLE, line=1.2)
footer(s, True)
notes(s, "Fill in the amounts honestly. Disclosure is not the safeguard; understanding is. Expect 'why velocity "
         "tracking and not kinematic?' and answer it yourselves.")

# 17 · Testing
s = clone(12, {"Image 0"})
header(s, "11 · Testing & evaluation", "Walkthrough tests by a non-builder", False)
text(s, L, 2.85, CW, 0.5, "<Who tested whose part · number of full runs · average time per shift>", size=19,
     color=MUTED)
tests = [("Launch", "Ticket readable, goal clear in 5 s"), ("Start", "START moves to 'unlock'"),
         ("Detach", "Handle swing frees portafilter"), ("Dose gating", "No grind unless in the fork"),
         ("Dose", "Counts to ~18 g; under 8 g refused"), ("Tamp", "Force rises, haptics, lift completes"),
         ("Tilt", "Tilted tamp gives Channeled"), ("Lock", "Brew refused until locked"),
         ("Brew", "Needle ~9 bar; time & yield count"), ("Grind effect", "Finer grind, longer shot"),
         ("Steam", "Heats only in milk; chime at 65 °C"), ("Pour", "Tilt fills cup; verdict appears"),
         ("Retry", "Restarts at unlock, grind kept"), ("Audio", "Each sound from its source"),
         ("Physics", "No floating or falling through"), ("Exported build", "Exe runs; shift done with mouse")]
tw = (CW - GUT) / 2
for half in range(2):
    rows = [["#", "Check", "Expected", "Result"]]
    for k in range(8):
        n = half * 8 + k
        rows.append([f"{n + 1:02d}", tests[n][0], tests[n][1], "☐ Pass  ☐ Fail"])
    table(s, L + half * (tw + GUT), TOP, [1.15, 2.75, tw - 1.15 - 2.75 - 2.45, 2.45], rows,
          row_h=(BOT - TOP) / 8.75, dark=False, size=17, head_size=14)
footer(s, False)
notes(s, "Tick results from real runs and add anything that failed and was fixed. Objectives 1-3 are measured "
         "here: time to finish, attempts to Balanced, final milk temperature.")

# 18 · Exported build check
s = clone(15, {"Image 0", "Image 2"})
header(s, "11 · Exported-build check", "Runs outside Unity, no headset", False)
photo(s, shot("exe_launch"), 2.39, 3.93, 10.82, 7.10, focus=(0.45, 0.5), rounded=False)
text(s, 2.39, 12.75, 10.82, 0.5, "Launching Builds/Windows/BaristaVR.exe after a successful build",
     size=17, color=MUTED, align=PP_ALIGN.CENTER)
facts = [("How it's built", "Barista ▸ Build Windows Player regenerates and saves the scene, then builds the exe."),
         ("How it's driven", "Keyboard and mouse bindings on the same XRI actions; the simulator isn't in a build."),
         ("What we saw", "Recorded run: launch to verdict in about 1.5 min. Shot 1 at 19.1 s → Sour; milk 65 °C."),
         ("Found & fixed", "The exe once showed an old scene, so the build now rebuilds the scene first.")]
for i, (h, b) in enumerate(facts):
    y = TOP + i * 2.35
    text(s, 15.6, y, R - 15.6, 2.2, [(h.upper(), {"size": 15, "font": FONT_SEMI, "spc": 200,
                                                  "color": CARAMEL_DEEP, "after": 6}),
                                     (b, {"size": 24, "color": INK})], line=1.2)
footer(s, False)
notes(s, "Named requirement in the brief. State exactly what you did: launched the exe without a headset and "
         "completed a shift with keyboard and mouse (see the recording).")

# 19 · Challenges
s = clone(2, {"Image 0", "Image 1", "Image 3"})
header(s, "12 · Challenges & solutions", "What went wrong, and how we fixed it", True)
chal = [("Unreadable floating text", "Labels blended into the café and showed through the hands.",
         "Custom WorldText shader with a hard edge, plus backdrop plates."),
        ("Exe showed an old scene", "The build used the last saved scene, not the generated one.",
         "The build menu now regenerates and saves the scene first."),
        ("Buttons pressed by accident", "Counter buttons sat in the working area.",
         "Moved to a back-wall station panel; Reset and Exit need two presses."),
        ("Floaty, jumpy tools", "Cups and pitchers had no weight and flew off.",
         "Capped velocity tracking, smoothing, 90 Hz physics, settle on release."),
        ("Tools passed through things", "Kinematic grabs ignored the counter and machine.",
         "Physical grabs with capped push-out; ignores only where parts nest."),
        ("Wobbly, sideways tamper", "It tilted in the hand and lay flat when tamping.",
         "Custom grab transformer: upright grip, limited wrist follow.")]
_, cw3 = grid_x(3, 0, 0.45)
inner = cw3 - 2 * PAD
ch = min((BOT - TOP - 0.45) / 2,
         max(2 * PAD + 0.62 + est_h(p, 22, inner) + 0.62 + est_h(f, 22, inner) for _, p, f in chal) + 0.2)
for i, (h, prob, fix) in enumerate(chal):
    x, w = grid_x(3, i % 3, 0.45)
    y = TOP + (i // 3) * (ch + 0.45)
    box(s, x, y, w, ch, fill=DARK_CARD)
    text(s, x + PAD, y + PAD, w - 2 * PAD, ch - 2 * PAD, [
        (h, {"size": 27, "font": FONT_SEMI, "color": CREAM, "after": 10}),
        (prob, {"size": 22, "color": MUTED_LIGHT, "after": 16}),
        ("FIX", {"size": 14, "font": FONT_SEMI, "spc": 200, "color": CARAMEL, "after": 4}),
        (fix, {"size": 22, "color": CREAM}),
    ], line=1.2)
footer(s, True)
notes(s, "Pick two or three you can explain best and go deep on why the fix works.")

# 20 · Limitations
s = clone(12, {"Image 0"})
header(s, "13 · Limitations & constraints", "What the simulator could not tell us", False)
table(s, L, TOP, [5.6, CW - 5.6], [
    ["Could not verify", "Published guidance or safe default used instead"],
    ["Comfort", "Seated, stationary, snap turn only, no artificial motion (Meta comfort guidance; LaViola, 2000)"],
    ["Scale & reach", "0.84 m counter, tools 0.4–0.7 m away, real dimensions (58 mm basket, 350 ml pitcher)"],
    ["Presence", "One coherent café, spatial audio, diegetic UI (Slater & Wilbur, 1997)"],
    ["Text legibility", "Panels at 0.75–1 m, 1.2–2.4 cm letters, high-contrast backdrops"],
    ["Haptics & force", "Haptics sent through XRI but not felt in the simulator; controllers can't push back"],
    ["Binaural audio", "Unity 3D panning; an HRTF spatializer is future work"],
], row_h=1.12, dark=False, size=20)
ly = TOP + 1.12 * 0.75 + 6 * 1.12 + 0.5
text(s, L, ly, CW, BOT - ly, [
    ("**Out of scope:** multiplayer, latte art, customer AI, real fluid simulation.", {"after": 8}),
    ("**Headset check:** <tested on a real headset? yes or no, and what you found>", {}),
], size=21, color=BODY, line=1.2)
footer(s, False)
notes(s, "The brief rewards stating what you could NOT verify and the guidance used instead. Don't overclaim.")

# 21 · Conclusion
s = clone(13, {"Image 0"})
header(s, "14 · Conclusion & reflection", "Did we solve the problem?", True)
lw = 11.0
res = [("Unaided, under 6 min", "<met / not met · average time>"),
       ("Balanced within 3 attempts", "<met / not met · attempts>"),
       ("Milk at 63–67 °C", "<met · 65 °C in the recorded run>")]
for i, (h, r_) in enumerate(res):
    y = TOP + i * 1.75
    box(s, L, y, lw, 1.5, fill=DARK_CARD)
    text(s, L + PAD, y, 1.0, 1.5, f"0{i + 1}", size=30, font=FONT_SEMI, color=CARAMEL, anchor=MSO_ANCHOR.MIDDLE)
    text(s, L + 1.6, y, lw - 1.6 - PAD, 1.5, [(h, {"size": 23, "font": FONT_SEMI, "color": CREAM, "after": 2}),
                                              (r_, {"size": 19, "color": MUTED_LIGHT})],
         anchor=MSO_ANCHOR.MIDDLE, line=1.1)
wy = TOP + 3 * 1.75 + 0.35
text(s, L, wy, lw, BOT - wy, [
    ("WHAT WE LEARNED", {"size": 15, "font": FONT_SEMI, "spc": 200, "color": CARAMEL, "after": 8}),
    ("<e.g. physical feel needs tuning, not just correct code; designing for a headset we couldn't always test "
     "meant leaning on published guidance>", {"size": 21, "color": CREAM}),
], line=1.22)
mx = L + lw + GUT + 0.3
mw = (R - mx - 0.45) / 2
mh = (BOT - TOP - 0.45) / 2
for i in range(4):
    x = mx + (i % 2) * (mw + 0.45)
    y = TOP + (i // 2) * (mh + 0.45)
    box(s, x, y, mw, mh, fill=DARK_CARD)
    text(s, x + PAD, y + PAD, mw - 2 * PAD, mh - 2 * PAD, [
        (f"<Member {i + 1}>", {"size": 23, "font": FONT_SEMI, "color": CREAM, "after": 10}),
        ("<One-minute reflection: what I built, what was hard, what I'd do differently>",
         {"size": 18, "color": MUTED_LIGHT}),
    ], line=1.2)
footer(s, True)
notes(s, "Answer the title with test numbers, then each member gives a short personal reflection (individual marks).")

# 22 · References
s = clone(12, {"Image 0"})
header(s, "15 · References & credits", "References", False)
left = ["Unity Technologies. XR Interaction Toolkit 3.6 manual. docs.unity3d.com",
        "Meta. VR best practices: locomotion, user comfort, interaction. developers.meta.com/horizon",
        "LaViola, J. J. (2000). A discussion of cybersickness in virtual environments. ACM SIGCHI Bulletin, 32(1).",
        "Hoffman, D. M., Girshick, A. R., Akeley, K., & Banks, M. S. (2008). Vergence-accommodation conflicts "
        "hinder visual performance and cause visual fatigue. Journal of Vision, 8(3).",
        "Slater, M., & Wilbur, S. (1997). A framework for immersive virtual environments (FIVE). Presence, 6(6).",
        "Specialty Coffee Association. Espresso brewing guidance (dose, yield, time)."]
right = ["XR Origin & XR Device Simulator: XR Interaction Toolkit 3.6.1 Starter Assets, Unity Companion License.",
         "Unity OpenXR Plugin, Input System, Universal Render Pipeline: Unity Technologies.",
         "Café models built by the team from Unity primitives; audio synthesised in code (AudioKit.cs).",
         "Build font: Unity LegacyRuntime (Arial). Slide font: Google Sans (SIL Open Font License).",
         "Slide template: “Coffee” free presentation template by EaTemp.",
         "Barista illustration: <source and licence>.",
         "AI tools: Cursor (see AI-tool disclosure)."]
for col, (head, items) in enumerate([("Guidance & theory", left), ("Assets & tools (also credited in the build)",
                                                                    right)]):
    x, w = grid_x(2, col, 1.0)
    text(s, x, TOP, w, 0.5, head.upper(), size=15, font=FONT_SEMI, spc=200, color=CARAMEL_DEEP)
    rule(s, x, TOP + 0.6, w)
    bullets(s, x, TOP + 0.9, w, BOT - TOP - 0.9, items, size=19, color=BODY, after=12, line=1.18)
footer(s, False)
notes(s, "Shown, not narrated. Fill in the barista illustration source before submitting.")

# Callback before the demo
s = dark_room()
line(s, 3.7, 3.0, "So.", size=80, color=MUTED_LIGHT, font=FONT)
line(s, 6.0, 3.0, "Would you like to brew a coffee?", size=88)
line(s, 9.3, 1.4, "Pull up a chair.", size=40, color=CARAMEL, font=FONT, italic=True)
page += 1
notes(s, "THE CALLBACK. Same speaker as the cold open. \"So.\" [Beat.] \"Would you like to brew a coffee?\" "
         "[Turn to the machine or headset.] \"Pull up a chair.\"\n\nThen straight into the live demo.")

# 23 · Live demo
s = clone(16)
fade(s, "med")
s.background.fill.solid()
s.background.fill.fore_color.rgb = DARK
px = 11.2
photo(s, shot("pour"), px, 1.15, R - px, BOT - 1.15, focus=(0.42, 0.5))
text(s, L, 1.15, 8.6, 0.45, "16 · LIVE DEMONSTRATION", size=17, font=FONT_SEMI, spc=250, color=CARAMEL)
text(s, L, 1.6, 8.6, 1.5, "Live demo", size=58, font=FONT_SEMI, color=CREAM, line=1.0)
text(s, L, 2.85, 8.6, 0.5, "Up to 10 minutes, including questions", size=21, color=MUTED_LIGHT)
demo = [("A", "Seated start, snap turn, no movement provider"), ("B", "Dose and tamp; tilt once to show channeling"),
        ("C", "Lock in, pull the shot, narrate gauge and model"),
        ("D", "Steam to 65 °C, pour, read the dial, pull another shot"),
        ("A", "Launch the exe and complete one action with the mouse")]
sy = TOP + 0.1
st = (BOT - sy) / 5
for i, (who, what) in enumerate(demo):
    y = sy + i * st
    box(s, L, y + (st - 1.05) / 2, 1.05, 1.05, fill=CARAMEL, shape=MSO_SHAPE.OVAL)
    text(s, L, y + (st - 1.05) / 2, 1.05, 1.05, who, size=26, font=FONT_SEMI, color=DARK, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, L + 1.5, y, px - L - 1.5 - GUT, st, what, size=22, color=CREAM, anchor=MSO_ANCHOR.MIDDLE, line=1.15)
footer(s, True)
notes(s, "Run live from START to the result dial without help, then each member operates their own part. Have the "
         "exe and the Editor both ready.")

# 24 · Thank you
s = clone(19, {"Image 0", "Image 1", "Shape 1"})
text(s, 3.92, 3.55, 10.6, 3.9, "Thank you.\nQuestions?", size=92, font=FONT_SEMI, color=CREAM, line=0.95)
rule(s, 3.95, 8.0, 1.4, color=CARAMEL, weight=3)
text(s, 3.92, 8.35, 10.6, 3.0, [
    ("TEAM XR CONTINUUM", {"size": 17, "font": FONT_SEMI, "spc": 250, "color": CARAMEL, "after": 10}),
    ("Barista Calibration Trainer · INTE 42312", {"size": 24, "color": CREAM, "after": 6}),
    ("<Member 1> · <Member 2> · <Member 3> · <Member 4>", {"size": 21, "color": MUTED_LIGHT}),
], line=1.2)
page += 1
notes(s, "Close and invite questions; keep the build running.")


# ------------------------------------------------------------------ finish

sld_ids = prs.slides._sldIdLst
for sld in list(sld_ids)[:len(TEMPLATE_SLIDES)]:
    prs.part.drop_rel(sld.get(qn("r:id")))
    sld_ids.remove(sld)

theme = prs.slide_master.part.part_related_by(RT.THEME)
xml = theme.blob.decode("utf-8")
xml = re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*(")', r"\g<1>%s\g<2>" % FONT, xml)
theme._blob = xml.encode("utf-8")

out = sys.argv[1] if len(sys.argv) > 1 else OUT
prs.save(out)
print("saved", out, "slides:", len(prs.slides))
